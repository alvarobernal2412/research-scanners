import socket

import pytest
from conftest import blank_pdf, make_pdf

from extract import ExtractError, extract, extract_bytes, split_lines


def _write(tmp_path, name, data):
    p = tmp_path / name
    p.write_bytes(data if isinstance(data, bytes) else data.encode("utf-8"))
    return p


def test_md_headings_and_blockquotes(tmp_path):
    doc = extract(_write(tmp_path, "a.md", "# Intro\nbody text\n> quoted words\nafter\n"))
    assert [l.heading_level for l in doc.lines] == [1, None, None, None]
    assert [l.in_quote for l in doc.lines] == [False, False, True, False]
    assert doc.lines[0].text == "Intro"
    assert doc.fmt == "md"


def test_md_fenced_code_blanked_line_numbers_kept(tmp_path):
    doc = extract(_write(tmp_path, "a.md", "before\n```py\ncode here\n```\nafter\n"))
    assert [l.no for l in doc.lines] == [1, 2, 3, 4, 5]
    assert [l.text for l in doc.lines] == ["before", "", "", "", "after"]


def test_crlf_cr_and_bom_line_numbers(tmp_path):
    doc = extract(_write(tmp_path, "a.txt", b"\xef\xbb\xbfone\r\ntwo\rthree\nfour"))
    assert [(l.no, l.text) for l in doc.lines] == [(1, "one"), (2, "two"), (3, "three"), (4, "four")]


def test_formfeed_does_not_split_lines(tmp_path):
    doc = extract(_write(tmp_path, "a.txt", "one\x0ctwo\nthree"))
    assert len(doc.lines) == 2
    assert split_lines("a\x0cb\r\nc") == ["a\x0cb", "c"]


def test_txt_passthrough(tmp_path):
    p = _write(tmp_path, "a.txt", "Hello world.\n\nSecond.\n")
    doc = extract(p)
    assert [l.text for l in doc.lines] == ["Hello world.", "", "Second."]
    assert len(doc.sha256) == 64


def test_docx_heading_styles_and_paragraphs(tmp_path):
    import docx

    d = docx.Document()
    d.add_heading("Top", level=1)
    d.add_paragraph("Plain paragraph.")
    d.add_heading("Sub", level=2)
    p = tmp_path / "a.docx"
    d.save(str(p))
    doc = extract(p)
    texts = [(l.text, l.heading_level) for l in doc.lines]
    assert ("Top", 1) in texts and ("Sub", 2) in texts and ("Plain paragraph.", None) in texts


def test_html_blocks_headings_blockquote_script_skipped(tmp_path):
    html = ("<html><head><style>x{}</style></head><body><h2>Methods</h2><p>First para.</p>"
            "<script>var secret=1;</script><blockquote>Quoted text</blockquote><p>Last</p></body></html>")
    doc = extract(_write(tmp_path, "a.html", html))
    by_text = {l.text: l for l in doc.lines}
    assert by_text["Methods"].heading_level == 2
    assert by_text["Quoted text"].in_quote and not by_text["First para."].in_quote
    assert not any("secret" in l.text or "x{}" in l.text for l in doc.lines)


def test_tex_section_quote_env_and_comments(tmp_path):
    tex = "\n".join([
        r"\section{Intro}",
        r"Text here 50\% sure. % hidden comment",
        r"\begin{quote}",
        "quoted line",
        r"\end{quote}",
        r"\subsection{Sub}",
        "plain",
        r"\begin{thebibliography}{9}",
        r"\bibitem{a} A. Author, Title, 2020.",
        r"\end{thebibliography}",
    ]) + "\n"
    doc = extract(_write(tmp_path, "a.tex", tex))
    lines = {l.no: l for l in doc.lines}
    assert lines[1].heading_level == 1 and lines[1].text == "Intro"
    assert r"50\% sure." in lines[2].text and "hidden" not in lines[2].text
    assert lines[4].in_quote and not lines[2].in_quote and not lines[7].in_quote
    assert lines[6].heading_level == 2
    assert "bibitem" in lines[9].text


def test_pdf_text_pages_and_ligature_nfkc(tmp_path):
    data = make_pdf(["Soﬀer et al."])
    doc = extract_bytes(data, "pdf")
    assert doc.lines[0].text == "Soffer et al."
    assert doc.lines[0].page == 1


def test_pdf_cid_artifacts_stripped():
    doc = extract_bytes(make_pdf(["Alpha (cid:123) beta (cid:4)gamma"]), "pdf")
    assert "cid" not in doc.lines[0].text


def test_pdf_pages_numbered():
    doc = extract_bytes(make_pdf(["one", "two", "three", "four"], pages=2), "pdf")
    assert sorted({l.page for l in doc.lines}) == [1, 2]


def test_scanned_pdf_raises_no_text_extracted(tmp_path):
    with pytest.raises(ExtractError) as e:
        extract(_write(tmp_path, "s.pdf", blank_pdf()))
    assert e.value.code == "no_text_extracted"


def test_empty_and_whitespace_txt_raise_no_text_extracted(tmp_path):
    for name, content in (("e.txt", ""), ("w.txt", " \n\t\n"), ("e.md", "")):
        with pytest.raises(ExtractError) as e:
            extract(_write(tmp_path, name, content))
        assert e.value.code == "no_text_extracted"


def test_unknown_extension_unsupported_format(tmp_path):
    with pytest.raises(ExtractError) as e:
        extract(_write(tmp_path, "a.xyz", "text"))
    assert e.value.code == "unsupported_format"


def test_missing_file_and_corrupt_pdf(tmp_path):
    with pytest.raises(ExtractError) as e:
        extract(tmp_path / "nope.txt")
    assert e.value.code == "file_not_found"
    with pytest.raises(ExtractError) as e:
        extract(_write(tmp_path, "bad.pdf", b"not a pdf at all"))
    assert e.value.code == "unreadable_file"


def test_extract_bytes_matches_extract_for_pdf(tmp_path):
    data = make_pdf(["Some words here.", "More words there."])
    p = _write(tmp_path, "a.pdf", data)
    a, b = extract(p), extract_bytes(data, "pdf")
    assert [(l.text, l.page) for l in a.lines] == [(l.text, l.page) for l in b.lines]
    assert a.sha256 == b.sha256


def test_network_blocked_in_tests():
    with pytest.raises(RuntimeError):
        socket.create_connection(("example.com", 80))

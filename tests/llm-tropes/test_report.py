import json
import subprocess
import sys
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parents[2] / "skills" / "llm-tropes" / "scripts"
sys.path.insert(0, str(SCRIPTS))

from patterns import load_patterns  # noqa: E402
from scan import build_report  # noqa: E402

DATA = SCRIPTS.parent / "data"
SCAN_PY = SCRIPTS / "scan.py"

LEX = "It was a perfect storm of factors."
STRUCT = "Not a bug. Not a feature. A flaw."
CHAT = "Would you like me to expand on this?"
FILLER = "The committee reviewed the budget for the coming year and approved the schedule."


def md(tmp_path, text, name="doc.md"):
    p = tmp_path / name
    p.write_text(text, encoding="utf-8")
    return p


def sec(report, i=0):
    return report["sections"][i]


def test_section_dense_when_three_categories(tmp_path):
    r = build_report(md(tmp_path, f"# A\n\n{LEX}\n\n{STRUCT}\n\n{CHAT}\n"), DATA)
    s = sec(r)
    assert s["heading"] == "# A" or s["heading"] == "A"
    assert len(s["categories"]) >= 3
    assert s["dense"] is True


def test_section_not_dense_with_two_categories(tmp_path):
    r = build_report(md(tmp_path, f"# A\n\n{LEX}\n\n{STRUCT}\n"), DATA)
    assert len(sec(r)["categories"]) == 2
    assert sec(r)["dense"] is False


def test_dense_ignores_contested_only_hits(tmp_path):
    r = build_report(md(tmp_path, "# A\n\nOne \u2014 two \u2014 three \u2014 four.\n"), DATA)
    s = sec(r)
    assert s["hit_count"] >= 3
    assert s["categories"] == []
    assert s["dense"] is False


def test_human_academic_fixture_not_dense(tmp_path):
    text = (
        "The survey of regional water districts covered forty municipalities over six years. "
        "However, response rates varied widely between the northern and southern counties. "
        "Moreover, the oldest records were kept on paper, so we digitised them by hand. "
        "Rainfall \u2014 which differed sharply from year to year \u2014 explained part of the variance in demand. "
        "Households with metered supply reduced consumption by roughly eight percent after the tariff change, "
        "whereas unmetered households showed no measurable response. "
        "Industrial users behaved differently \u2014 their consumption tracked production schedules rather than price. "
        "We therefore analyse the two groups separately and report confidence intervals for each estimate. "
        "The limitations of the data are discussed in the final section, together with directions for further work."
    )
    r = build_report(md(tmp_path, text), DATA, "academic")
    assert all(not s["dense"] for s in r["sections"])
    assert r["stats"]["words"] > 100


def test_all_code_markdown_zero_findings_exit_0(tmp_path):
    p = md(tmp_path, "```\nperfect storm\n```\n")
    res = subprocess.run([sys.executable, "-I", str(SCAN_PY), str(p)], capture_output=True)
    assert res.returncode == 0
    out = json.loads(res.stdout.decode("utf-8"))
    assert out["findings"] == []
    assert out["error"] is None


def test_non_english_text_sets_language_note_no_findings(tmp_path):
    text = (
        "El comit\u00e9 revis\u00f3 el presupuesto anual y aprob\u00f3 el calendario de actividades para los pr\u00f3ximos meses. "
        "Los participantes discutieron las propuestas con calma, y despu\u00e9s votaron por mayor\u00eda."
    )
    r = build_report(md(tmp_path, text), DATA)
    assert r["language_note"]
    assert r["findings"] == []


def test_cli_prints_valid_json_with_disclaimer(tmp_path):
    p = md(tmp_path, f"{FILLER}\n\n{LEX}\n")
    res = subprocess.run([sys.executable, "-I", str(SCAN_PY), str(p)], capture_output=True)
    assert res.returncode == 0
    out = json.loads(res.stdout.decode("utf-8"))
    for k in ("file", "profile", "findings", "sections", "stats", "density_per_1k",
              "language_note", "disclaimer", "error"):
        assert k in out
    assert "authorship" in out["disclaimer"]
    assert out["density_per_1k"] > 0


def _blank_pdf() -> bytes:
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode() + b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


def test_cli_no_text_pdf_exits_2_with_no_text_extracted(tmp_path):
    p = tmp_path / "blank.pdf"
    p.write_bytes(_blank_pdf())
    res = subprocess.run([sys.executable, "-I", str(SCAN_PY), str(p)], capture_output=True)
    assert res.returncode == 2
    assert json.loads(res.stdout.decode("utf-8")) == {"error": "no_text_extracted"}


def test_cli_malformed_pdf_exits_2_with_unreadable_file(tmp_path):
    p = tmp_path / "bad.pdf"
    p.write_bytes(b"%PDF-1.4\n%%EOF\n")
    res = subprocess.run([sys.executable, "-I", str(SCAN_PY), str(p)], capture_output=True)
    assert res.returncode == 2
    assert json.loads(res.stdout.decode("utf-8")) == {"error": "unreadable_file"}


def test_profile_filters_patterns(tmp_path):
    d = tmp_path / "pats"
    d.mkdir()
    base = {"category": "lexical", "weight": 1.0, "source": "t", "explanation": "x", "example": "x"}
    (d / "p.json").write_text(json.dumps([
        {**base, "id": "all", "regex": r"\bfoo\b"},
        {**base, "id": "acad", "regex": r"\bbar\b", "profiles": ["academic"]},
    ]), encoding="utf-8")
    f = md(tmp_path, "foo and bar here.\n")
    ids = lambda prof: {x["pattern_id"] for x in build_report(f, d, prof)["findings"]}
    assert ids("academic") == {"all", "acad"}
    assert ids("general") == {"all"}


# ---- final-review regressions -------------------------------------------------

from scan import DENSE_MIN_PER_1K, find_hits as _find_hits  # noqa: E402
from extract import Line as _Line  # noqa: E402


def _real_hits(text):
    return _find_hits([_Line(1, text)], load_patterns(DATA))


_PLAIN = (
    "The survey of regional water districts covered forty municipalities over six years. "
    "Response rates varied between the northern and southern counties, and older records were kept on paper. "
    "Households with metered supply reduced consumption after the tariff change, while others showed little response. "
    "We analyse the two groups separately and report confidence intervals for each estimate. "
)


def test_c1_long_single_section_txt_with_three_phrases_not_dense(tmp_path):
    text = (_PLAIN * 60) + f"\n{LEX}\n{STRUCT}\n{CHAT}\n" + (_PLAIN * 20)
    p = tmp_path / "paper.txt"
    p.write_text(text, encoding="utf-8")
    r = build_report(p, DATA)
    s = r["sections"][0]
    assert len(s["categories"]) >= 3
    assert s["density_per_1k"] < DENSE_MIN_PER_1K
    assert s["dense"] is False


def test_c1_short_slop_paragraph_dense_and_has_density_field(tmp_path):
    r = build_report(md(tmp_path, f"{LEX} {STRUCT} {CHAT}\n"), DATA)
    s = r["sections"][0]
    assert s["dense"] is True
    assert s["density_per_1k"] >= DENSE_MIN_PER_1K


def test_c1_multi_category_human_text_not_dense(tmp_path):
    nl = chr(10) * 2
    text = nl.join([
        (_PLAIN * 5) + "Moreover, the oldest records were kept on paper.",
        (_PLAIN * 5) + "Water pricing plays a crucial role in demand management.",
        (_PLAIN * 5) + "In conclusion, we recommend further study.",
        (_PLAIN * 5) + "The 2015 reform marks a significant shift in national policy.",
    ])
    r = build_report(md(tmp_path, text), DATA)
    s = r["sections"][0]
    assert len(s["categories"]) >= 3
    assert s["density_per_1k"] < DENSE_MIN_PER_1K / 2
    assert s["dense"] is False


def test_i1_overlapping_entries_keep_one_finding():
    for phrase in ("fundamentally reshape", "ever-evolving landscape", "pivotal role"):
        hits = _real_hits(f"This will {phrase} things.")
        assert len(hits) == 1, (phrase, [h.pattern_id for h in hits])


def test_i1_overlap_does_not_make_sentence_dense(tmp_path):
    r = build_report(md(tmp_path, "Moreover, this will fundamentally reshape the field.\n"), DATA)
    assert len(r["sections"][0]["categories"]) < 3
    assert r["sections"][0]["dense"] is False


def test_i1_non_overlapping_findings_untouched():
    hits = _real_hits("Moreover, it was a perfect storm.")
    assert len(hits) >= 2


def test_i3_typographic_apostrophes_flagged():
    for s in ("Here’s the kicker", "Let’s break this down", "You’re absolutely right!",
              "It’s not a bug. It’s a feature."):
        hits = _real_hits(s)
        assert hits, s
        for h in hits:
            assert h.text in s  # original text slice, curly apostrophes preserved


def test_i3_original_text_reported():
    hits = _real_hits("Here’s the kicker")
    assert any("’" in h.text for h in hits)


def test_i6_contested_excluded_from_density(tmp_path):
    p = md(tmp_path, "One — two — three — four — five “quoted” words here for the test.\n")
    r = build_report(p, DATA)
    assert r["findings"] and all(f["contested"] for f in r["findings"])
    assert r["density_per_1k"] == 0.0


def _cli(*args):
    return subprocess.run([sys.executable, "-I", str(SCAN_PY), *args], capture_output=True)


def test_i2_nonexistent_patterns_dir(tmp_path):
    p = md(tmp_path, "hello world\n")
    res = _cli(str(p), "--patterns-dir", str(tmp_path / "nowhere"))
    assert res.returncode == 2
    out = json.loads(res.stdout.decode("utf-8"))
    assert out["error"] == "catalog_error" and out["detail"]


def test_i2_malformed_json_in_patterns_dir(tmp_path):
    d = tmp_path / "pats"
    d.mkdir()
    (d / "p.json").write_text("{not json", encoding="utf-8")
    res = _cli(str(md(tmp_path, "hello\n")), "--patterns-dir", str(d))
    assert res.returncode == 2
    assert json.loads(res.stdout.decode("utf-8"))["error"] == "catalog_error"


def test_i2_empty_patterns_dir(tmp_path):
    d = tmp_path / "pats"
    d.mkdir()
    res = _cli(str(md(tmp_path, "hello\n")), "--patterns-dir", str(d))
    assert res.returncode == 2
    assert json.loads(res.stdout.decode("utf-8"))["error"] == "catalog_error"


def test_m3_empty_and_whitespace_only_input(tmp_path):
    for content in ("", "   \n\n \t\n"):
        res = _cli(str(md(tmp_path, content)))
        assert res.returncode == 2, repr(content)
        assert json.loads(res.stdout.decode("utf-8")) == {"error": "no_text_extracted"}


def test_m3_docx_without_paragraph_text(tmp_path):
    import docx

    d = docx.Document()
    d.add_table(rows=1, cols=1).cell(0, 0).text = "only a table"
    p = tmp_path / "t.docx"
    d.save(str(p))
    res = _cli(str(p))
    assert res.returncode == 2
    assert json.loads(res.stdout.decode("utf-8")) == {"error": "no_text_extracted"}

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "skills" / "llm-tropes" / "scripts"))

from extract import ExtractError, Line, extract  # noqa: E402


def test_md_marks_fenced_code_lines(tmp_path):
    p = tmp_path / "a.md"
    p.write_text("before\n```python\ncode here\n```\nafter\n", encoding="utf-8")
    lines = extract(p)
    assert [l.in_code for l in lines] == [False, True, True, True, False]
    assert lines[4].text == "after"


def test_md_marks_blockquote_lines(tmp_path):
    p = tmp_path / "a.md"
    p.write_text("plain\n> quoted\n>also quoted\nplain again\n", encoding="utf-8")
    lines = extract(p)
    assert [l.in_quote for l in lines] == [False, True, True, False]
    assert all(not l.in_code for l in lines)


def test_crlf_and_bom_line_numbers(tmp_path):
    p = tmp_path / "a.md"
    p.write_bytes(b"\xef\xbb\xbfone\r\ntwo\r\nthree\r\n")
    lines = extract(p)
    assert [(l.no, l.text) for l in lines] == [(1, "one"), (2, "two"), (3, "three")]


def test_txt_passthrough(tmp_path):
    p = tmp_path / "a.txt"
    p.write_text("```\nnot code\n```\n> not quote\n", encoding="utf-8")
    lines = extract(p)
    assert [l.text for l in lines] == ["```", "not code", "```", "> not quote"]
    assert not any(l.in_code or l.in_quote for l in lines)


def test_docx_paragraphs(tmp_path):
    import docx

    d = docx.Document()
    d.add_paragraph("First paragraph.")
    d.add_paragraph("Second paragraph.")
    p = tmp_path / "a.docx"
    d.save(str(p))
    lines = extract(p)
    assert [l.text for l in lines] == ["First paragraph.", "Second paragraph."]
    assert [l.no for l in lines] == [1, 2]


def _blank_pdf() -> bytes:
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode() + b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


def test_pdf_without_text_raises_no_text_extracted(tmp_path):
    p = tmp_path / "a.pdf"
    p.write_bytes(_blank_pdf())
    with pytest.raises(ExtractError) as ei:
        extract(p)
    assert ei.value.code == "no_text_extracted"


def test_unknown_extension_raises_unsupported_format(tmp_path):
    p = tmp_path / "a.xyz"
    p.write_text("hi")
    with pytest.raises(ExtractError) as ei:
        extract(p)
    assert ei.value.code == "unsupported_format"


def test_i5_formfeed_does_not_shift_line_numbers(tmp_path):
    p = tmp_path / "a.txt"
    p.write_bytes("a\x0cb\nIt was a perfect storm.".encode("utf-8"))
    lines = extract(p)
    assert [l.no for l in lines] == [1, 2]
    assert lines[1].text == "It was a perfect storm."


def test_i5_unicode_line_separator_and_lone_cr(tmp_path):
    p = tmp_path / "a.txt"
    p.write_bytes("a b\nsecond".encode("utf-8"))
    assert [(l.no, l.text) for l in extract(p)] == [(1, "a b"), (2, "second")]
    q = tmp_path / "b.txt"
    q.write_bytes(b"one\rtwo\rthree")
    assert [(l.no, l.text) for l in extract(q)] == [(1, "one"), (2, "two"), (3, "three")]


def test_i5_trailing_newline_adds_no_empty_line(tmp_path):
    p = tmp_path / "a.txt"
    p.write_bytes(b"x\ny\n")
    assert [l.text for l in extract(p)] == ["x", "y"]
